"""
DOCX document parser using python-docx.
"""

import io
from typing import Any, BinaryIO

from docx import Document

from app.core.errors import DocumentProcessingError
from app.core.logging import get_logger
from app.ingestion.parsers.base import BaseParser, ParsedDocument

logger = get_logger(__name__)


class DocxParser(BaseParser):
    """Parser for Microsoft Word DOCX documents."""
    
    @property
    def supported_extensions(self) -> list[str]:
        return ['.docx']
    
    @property
    def mime_types(self) -> list[str]:
        return [
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            'application/msword',  # Sometimes used for .docx
        ]
    
    async def parse_file(self, file_content: bytes, filename: str) -> ParsedDocument:
        """Parse DOCX from bytes."""
        try:
            file_stream = io.BytesIO(file_content)
            return await self.parse_stream(file_stream, filename)
        except Exception as e:
            logger.error("DOCX parsing failed", filename=filename, error=str(e))
            raise DocumentProcessingError(f"Failed to parse DOCX: {str(e)}")
    
    async def parse_stream(self, file_stream: BinaryIO, filename: str) -> ParsedDocument:
        """Parse DOCX from stream."""
        try:
            # Read DOCX document
            doc = Document(file_stream)
            
            # Extract text from paragraphs
            paragraphs = []
            paragraph_texts = []
            
            for para_idx, paragraph in enumerate(doc.paragraphs):
                text = paragraph.text.strip()
                if text:
                    paragraph_texts.append(text)
                    
                    # Extract paragraph metadata
                    para_info = {
                        "paragraph_index": para_idx,
                        "text": text,
                        "char_count": len(text),
                    }
                    
                    # Try to extract style information
                    try:
                        if paragraph.style:
                            para_info["style"] = paragraph.style.name
                        
                        # Check if it looks like a heading
                        if paragraph.style and any(
                            keyword in paragraph.style.name.lower() 
                            for keyword in ['heading', 'title', 'header']
                        ):
                            para_info["is_heading"] = True
                            para_info["heading_level"] = self._extract_heading_level(paragraph.style.name)
                        
                    except Exception as e:
                        logger.debug(f"Could not extract paragraph style: {e}")
                    
                    paragraphs.append(para_info)
            
            if not paragraph_texts:
                raise DocumentProcessingError("No text content found in DOCX")
            
            # Combine all paragraphs
            content = "\n\n".join(paragraph_texts)
            content = self._validate_content(content)
            
            # Build metadata
            metadata = self._extract_basic_metadata(filename, len(content.encode('utf-8')))
            metadata.update({
                "total_paragraphs": len(doc.paragraphs),
                "paragraphs_with_text": len(paragraphs),
                "paragraphs": paragraphs,
                "document_properties": self._extract_docx_properties(doc),
                "tables_count": len(doc.tables),
                "sections_count": len(doc.sections),
            })
            
            # Extract table content if present
            if doc.tables:
                metadata["tables"] = self._extract_tables(doc.tables)
            
            logger.info(
                "DOCX parsed successfully",
                filename=filename,
                paragraphs=len(doc.paragraphs),
                text_paragraphs=len(paragraphs),
                content_chars=len(content),
                tables=len(doc.tables),
            )
            
            return ParsedDocument.create(content, metadata)
            
        except DocumentProcessingError:
            raise
        except Exception as e:
            logger.error("DOCX parsing failed", filename=filename, error=str(e))
            raise DocumentProcessingError(f"Failed to parse DOCX: {str(e)}")
    
    def _extract_docx_properties(self, doc: Document) -> dict[str, Any]:
        """Extract DOCX document properties."""
        properties = {}
        
        try:
            core_props = doc.core_properties
            
            # Common document properties
            prop_mapping = {
                'title': 'title',
                'author': 'author', 
                'subject': 'subject',
                'keywords': 'keywords',
                'category': 'category',
                'comments': 'comments',
                'created': 'created',
                'modified': 'modified',
                'last_modified_by': 'last_modified_by',
                'revision': 'revision',
            }
            
            for prop_name, clean_name in prop_mapping.items():
                try:
                    value = getattr(core_props, prop_name, None)
                    if value:
                        if hasattr(value, 'isoformat'):  # DateTime object
                            properties[clean_name] = value.isoformat()
                        else:
                            properties[clean_name] = str(value)
                except Exception:
                    continue
                    
        except Exception as e:
            logger.warning("Failed to extract DOCX properties", error=str(e))
        
        return properties
    
    def _extract_heading_level(self, style_name: str) -> int:
        """Extract heading level from style name."""
        style_lower = style_name.lower()
        
        # Look for numbers in heading styles
        for i in range(1, 7):  # Heading 1 through 6
            if f"heading {i}" in style_lower or f"heading{i}" in style_lower:
                return i
        
        # Default to level 1 for unrecognized heading styles
        return 1
    
    def _extract_tables(self, tables: list) -> list[dict[str, Any]]:
        """Extract table content and structure."""
        table_data = []
        
        for table_idx, table in enumerate(tables):
            try:
                rows = []
                
                for row_idx, row in enumerate(table.rows):
                    cells = []
                    for cell_idx, cell in enumerate(row.cells):
                        cell_text = cell.text.strip()
                        if cell_text:
                            cells.append({
                                "column": cell_idx,
                                "text": cell_text,
                            })
                    
                    if cells:
                        rows.append({
                            "row": row_idx,
                            "cells": cells,
                        })
                
                if rows:
                    table_info = {
                        "table_index": table_idx,
                        "rows": rows,
                        "total_rows": len(table.rows),
                        "total_columns": len(table.columns) if table.rows else 0,
                    }
                    table_data.append(table_info)
                    
            except Exception as e:
                logger.warning(f"Failed to extract table {table_idx}: {e}")
                continue
        
        return table_data

